"""Тесты выгрузок .docx (программа, отчёт) и .pptx (слайды) — karkas.infra.exporters.

Фикстура комплекта — ``build_export_kit()`` из test_exporters_gift.py; минимальный
комплект — фикстура ``minimal_kit`` из tests/conftest.py.
"""

from __future__ import annotations

import io

import pytest
from docx import Document
from docx.document import Document as DocxDocument
from docx.table import Table
from pptx import Presentation

from karkas.domain.kit import ContentBlock, KitV1, Report, ReportIssue
from karkas.infra.exporters.docx_program import DOP_ACR_SECTION_HEADINGS, export_docx_program
from karkas.infra.exporters.docx_report import export_docx_report
from karkas.infra.exporters.pptx_slides import export_pptx_slides
from tests.conftest import build_minimal_kit
from tests.unit.test_exporters_gift import build_export_kit

TOTAL_COLUMN = "Всего, ч"
UNVERIFIED_MARK = "[не подтверждено — проверьте]"


@pytest.fixture()
def kit() -> KitV1:
    return build_export_kit()


def _docx_text(document: DocxDocument) -> str:
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def _find_plan_table(document: DocxDocument) -> Table:
    """Первая таблица с заголовком учебного плана (2.1)."""
    for table in document.tables:
        header = [cell.text for cell in table.rows[0].cells]
        if TOTAL_COLUMN in header and "ЛК" in header:
            return table
    raise AssertionError("таблица учебного плана не найдена")


def _pptx_text(presentation: Presentation) -> str:
    parts: list[str] = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
        if slide.has_notes_slide:
            parts.append(slide.notes_slide.notes_text_frame.text)
    return "\n".join(parts)


# ---------------------------------------------------------------------------
# Программа .docx
# ---------------------------------------------------------------------------


def test_docx_dop_follows_acr_section_headings() -> None:
    """ДОП: заголовки как в эталоне АЦР (data/templates/acr_dop_intensive_ai_24h)."""
    kit = build_minimal_kit(total_hours=24, final_att_hours=2)
    kit.program.title = "Интенсив по созданию продуктов на основе ИИ"
    text = _docx_text(Document(io.BytesIO(export_docx_program(kit))))
    missing = [heading for heading in DOP_ACR_SECTION_HEADINGS if heading not in text]
    assert not missing, "нет заголовков: " + ", ".join(missing)
    assert "Приложение 1. Оценочные материалы текущего контроля" not in text


def test_docx_program_is_zip_and_opens(kit: KitV1) -> None:
    data = export_docx_program(kit)
    assert data[:2] == b"PK"
    document = Document(io.BytesIO(data))
    text = _docx_text(document)
    assert kit.program.title in text
    assert "Учебный план" in text
    assert "М2 Обработка данных в pandas" in text
    assert "Аннотация" in text
    assert "Приложение 1. Оценочные материалы текущего контроля" in text
    assert "Приложение 2. Оценочные материалы итоговой аттестации" in text


def test_docx_program_totals_computed_from_topics(kit: KitV1) -> None:
    """Σ тем + ИА = объём: строки модулей и ИТОГО равны вычисленным значениям."""
    document = Document(io.BytesIO(export_docx_program(kit)))
    table = _find_plan_table(document)
    header = [cell.text for cell in table.rows[0].cells]
    total_index = header.index(TOTAL_COLUMN)
    rows = {row.cells[1].text: row for row in table.rows[1:]}
    module_1 = kit.framework.modules[0]
    assert int(rows["М1 Основы Python"].cells[total_index].text) == module_1.workload.total
    assert int(rows["М2 Обработка данных в pandas"].cells[total_index].text) == 14
    assert int(rows["Итоговая аттестация"].cells[total_index].text) == 2
    assert int(rows["ИТОГО"].cells[total_index].text) == kit.grand_total()
    assert kit.grand_total() == kit.program.total_hours == 36


def test_docx_program_theme_plan_contains_topics_and_distance_column(kit: KitV1) -> None:
    document = Document(io.BytesIO(export_docx_program(kit)))
    tables = [
        table
        for table in document.tables
        if TOTAL_COLUMN in [cell.text for cell in table.rows[0].cells]
    ]
    assert len(tables) == 2  # 2.1 учебный план и 2.2 учебно-тематический план
    thematic = tables[1]
    header = [cell.text for cell in thematic.rows[0].cells]
    assert "ДОТ" in header  # у тем есть дистанционные часы
    names = [row.cells[1].text for row in thematic.rows[1:]]
    assert "Введение в Python" in names
    assert "Функции и модули" in names


def test_docx_program_has_blank_order_and_signature_fields(kit: KitV1) -> None:
    """Реквизиты приказа и подписи — только пропуски-подчёркивания (V-TPL-03)."""
    document = Document(io.BytesIO(export_docx_program(kit)))
    text = _docx_text(document)
    assert "Приказом № ___ от" in text
    assert "Разработчик(и): ______" in text
    assert kit.placeholders.order_number == ""


def test_docx_program_marks_unverified_source(kit: KitV1) -> None:
    document = Document(io.BytesIO(export_docx_program(kit)))
    text = _docx_text(document)
    assert UNVERIFIED_MARK in text
    unverified_line = next(line for line in text.splitlines() if "Иванов И.И." in line)
    assert UNVERIFIED_MARK in unverified_line
    verified_line = next(line for line in text.splitlines() if "Python 3.12" in line)
    assert UNVERIFIED_MARK not in verified_line


def test_docx_program_uses_docxtpl_template(tmp_path) -> None:
    template = Document()
    template.add_paragraph("Программа: {{ title }}")
    template.add_paragraph("Объём: {{ total_hours }} ч")
    template.add_paragraph("Организация: {{ organization }}")
    path = tmp_path / "template.docx"
    template.save(path)

    kit = build_export_kit()
    data = export_docx_program(kit, template_path=path)
    text = _docx_text(Document(io.BytesIO(data)))
    assert f"Программа: {kit.program.title}" in text
    assert "Объём: 36 ч" in text
    assert "Организация: ООО «АЦР»" in text


def test_docx_program_missing_template_falls_back_to_generation(tmp_path) -> None:
    kit = build_export_kit()
    data = export_docx_program(kit, template_path=tmp_path / "absent.docx")
    text = _docx_text(Document(io.BytesIO(data)))
    assert "Учебный план" in text


def test_docx_program_template_is_sandboxed(tmp_path) -> None:
    """Шаблон рендерится в песочнице: доступ к опасным атрибутам не проходит."""
    template = Document()
    template.add_paragraph("Класс: {{ title.__class__ }}")
    template.add_paragraph("Название: {{ title }}")
    path = tmp_path / "sandbox.docx"
    template.save(path)
    data = export_docx_program(build_export_kit(), template_path=path)
    text = _docx_text(Document(io.BytesIO(data)))
    assert "Название: Анализ данных на Python" in text
    assert "class 'str'" not in text
    assert "__class__" not in text


def test_docx_program_minimal_kit_does_not_fail(minimal_kit: KitV1) -> None:
    """Пустой комплект (без тестов, слайдов, источников) не роняет экспорт."""
    data = export_docx_program(minimal_kit)
    assert data[:2] == b"PK"
    text = _docx_text(Document(io.BytesIO(data)))
    assert minimal_kit.program.title in text
    assert "Учебный план" in text


def test_docx_program_volume_line_uses_computed_hours(kit: KitV1) -> None:
    """1.8: контактная/СРС/ИА — из тем (26 + 4 + 4 + ИА 2 = 36)."""
    text = _docx_text(Document(io.BytesIO(export_docx_program(kit))))
    assert "Общая трудоёмкость — 36 акад. ч" in text
    assert "контактная работа — 26 ч" in text
    assert "текущий контроль — 4 ч" in text
    assert "самостоятельная работа — 4 ч" in text
    assert "итоговая аттестация — 2 ч (защита итогового проекта)" in text


# ---------------------------------------------------------------------------
# Отчёт .docx
# ---------------------------------------------------------------------------


def test_docx_report_contains_issue_assumptions_and_gaps(kit: KitV1) -> None:
    data = export_docx_report(kit)  # берётся kit.report
    assert data[:2] == b"PK"
    text = _docx_text(Document(io.BytesIO(data)))
    assert "V-SRC-04" in text
    assert "предупреждение" in text
    assert "Книга не найдена в доступных каталогах" in text
    assert "SRC-02" in text
    assert "Режим занятий не указан" in text
    assert "Образец: итоговая строка СРС" in text
    assert "уровень senior" in text
    assert "Проверка не пройдена" in text
    assert "версия комплекта: 3" in text


def test_docx_report_accepts_explicit_report(kit: KitV1) -> None:
    report = Report(
        passed=True,
        duration_ms=10,
        summary={"hours": {"ok": True, "checked": 3}},
        issues=[
            ReportIssue(
                rule_code="V-HRS-01",
                severity="error",
                message="Σ строк + ИА ≠ объём",
            )
        ],
    )
    text = _docx_text(Document(io.BytesIO(export_docx_report(kit, report=report))))
    assert "V-HRS-01" in text
    assert "ошибка" in text
    assert "Проверка пройдена" in text


def test_docx_report_without_report_gives_explicit_note() -> None:
    kit = build_export_kit()
    kit.report = None
    text = _docx_text(Document(io.BytesIO(export_docx_report(kit))))
    assert "Отчёт не сформирован" in text


# ---------------------------------------------------------------------------
# Слайды .pptx
# ---------------------------------------------------------------------------


def test_pptx_slides_generated_for_module_with_slides(kit: KitV1) -> None:
    data = export_pptx_slides(kit)
    assert data[:2] == b"PK"
    presentation = Presentation(io.BytesIO(data))
    assert len(presentation.slides) >= 2
    text = _pptx_text(presentation)
    assert "Основы Python" in text
    assert "Введение в Python" in text
    assert "План лекции" in text  # заметки докладчика из lecture_plan
    assert "Напишите программу" in text  # практика из задания семинара


def test_pptx_module_refs_filter(kit: KitV1) -> None:
    presentation = Presentation(io.BytesIO(export_pptx_slides(kit, module_refs=["M2"])))
    text = _pptx_text(presentation)
    assert "М2 Обработка данных в pandas" in text
    assert "Основы Python" not in text
    assert "pandas — это библиотека" in text  # вопрос модуля M2


def test_pptx_without_slides_returns_minimal_title_slide(minimal_kit: KitV1) -> None:
    data = export_pptx_slides(minimal_kit)
    assert data[:2] == b"PK"
    presentation = Presentation(io.BytesIO(data))
    assert len(presentation.slides) >= 1
    assert minimal_kit.program.title in _pptx_text(presentation)


def test_pptx_unknown_module_ref_returns_minimal_slide(kit: KitV1) -> None:
    """Несуществующий module_ref не роняет функцию: минимальный титульный слайд."""
    presentation = Presentation(io.BytesIO(export_pptx_slides(kit, module_refs=["M99"])))
    assert len(presentation.slides) >= 1
    assert kit.program.title in _pptx_text(presentation)


def test_pptx_nested_slides_structure(kit: KitV1) -> None:
    """body.slides (список слайдов) разворачивается в отдельные слайды."""
    kit.content.slides = [
        ContentBlock(
            ref="SL-T2.1-1",
            kind="slides",
            module_ref="M2",
            topic_ref="T2.1",
            body={
                "slides": [
                    {"title": "Первый слайд", "bullets": ["раз", "два"]},
                    {"title": "Второй слайд", "lines": "три\nчетыре"},
                ]
            },
        )
    ]
    presentation = Presentation(io.BytesIO(export_pptx_slides(kit)))
    text = _pptx_text(presentation)
    assert "Первый слайд" in text
    assert "Второй слайд" in text
    assert len(presentation.slides) >= 4


def test_pptx_content_slide_falls_back_to_topic_title(kit: KitV1) -> None:
    """Без title в body заголовок слайда — название темы из topic_ref."""
    kit.content.slides = [
        ContentBlock(
            ref="SL-T2.1-2",
            kind="slides",
            module_ref="M2",
            topic_ref="T2.1",
            body={"bullets": ["факт"]},
        )
    ]
    text = _pptx_text(Presentation(io.BytesIO(export_pptx_slides(kit))))
    assert "Анализ данных в pandas" in text
