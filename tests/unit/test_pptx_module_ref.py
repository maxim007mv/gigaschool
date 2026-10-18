"""PPTX: слайды с topic_ref без module_ref попадают в выгрузку."""

from __future__ import annotations

import io

from pptx import Presentation

from karkas.domain.kit import ContentBlock
from karkas.eval.reference_kits import build_tk1_kit
from karkas.infra.exporters.pptx_slides import export_pptx_slides


def test_pptx_slide_count_matches_blocks() -> None:
    kit = build_tk1_kit()
    kit.content.slides = [
        ContentBlock(
            ref="SL-T1.1",
            kind="slides",
            topic_ref="T1.1",
            module_ref=None,
            body={
                "slides": [
                    {"title": "Слайд 1", "bullets": ["a"], "speaker_notes": "n"},
                    {"title": "Слайд 2", "bullets": ["b"], "speaker_notes": "n"},
                ]
            },
        )
    ]
    data = export_pptx_slides(kit)
    presentation = Presentation(io.BytesIO(data))
    assert len(presentation.slides) == 2
