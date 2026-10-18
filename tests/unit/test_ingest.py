"""Тесты ingest: извлечение текста .docx/.pdf/.txt, проверки MIME и размера."""

from __future__ import annotations

import io

import pytest

from karkas.api.errors import FileTooLarge, UnsupportedFile
from karkas.infra import ingest


def _make_docx_bytes() -> bytes:
    from docx import Document

    document = Document()
    document.add_heading("Учебный план", level=1)
    document.add_paragraph("Программа повышения квалификации на 72 часа.")
    table = document.add_table(rows=2, cols=3)
    table.cell(0, 0).text = "Тема"
    table.cell(0, 1).text = "Лекции"
    table.cell(0, 2).text = "Практика"
    table.cell(1, 0).text = "T1.1 Введение"
    table.cell(1, 1).text = "2"
    table.cell(1, 2).text = "4"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_extract_docx_keeps_tables_and_headings() -> None:
    data = _make_docx_bytes()
    mime, text, meta = ingest.extract("plan.docx", data)
    assert mime.endswith("wordprocessingml.document")
    assert "Учебный план" in text
    assert "T1.1 Введение | 2 | 4" in text
    assert meta["tables"] == 1
    assert "Учебный план" in meta["headings"]


def test_extract_txt_and_md() -> None:
    mime, text, meta = ingest.extract("notes.txt", "Бриф: курс на 72 часа".encode())
    assert mime == "text/plain"
    assert "72" in text
    assert meta["lines"] == 1

    mime_md, text_md, _ = ingest.extract("readme.md", "# Заголовок\nтекст".encode())
    assert mime_md == "text/markdown"
    assert "Заголовок" in text_md


def test_extract_txt_cp1251_fallback() -> None:
    _, text, meta = ingest.extract("old.txt", "Пример".encode("cp1251"))
    assert text == "Пример"
    assert meta["encoding"] == "cp1251"


def test_unsupported_extension() -> None:
    with pytest.raises(UnsupportedFile):
        ingest.extract("script.exe", b"MZ\x90\x00")


def test_fake_pdf_signature_rejected() -> None:
    with pytest.raises(UnsupportedFile):
        ingest.extract("fake.pdf", b"not a pdf at all")


def test_fake_docx_signature_rejected() -> None:
    with pytest.raises(UnsupportedFile):
        ingest.extract("fake.docx", b"plain text pretending to be docx")


def test_binary_txt_rejected() -> None:
    with pytest.raises(UnsupportedFile):
        ingest.extract("data.txt", b"\x00\x01\x02binary")


def test_size_limits() -> None:
    with pytest.raises(FileTooLarge):
        ingest.extract("empty.txt", b"")
    too_big = b"a" * (ingest.settings.file_max_bytes + 1)
    with pytest.raises(FileTooLarge):
        ingest.extract("big.txt", too_big)


def test_control_chars_stripped() -> None:
    _, text, _ = ingest.extract("weird.txt", "текст\x07с\x08мусором".encode())
    assert "\x07" not in text and "\x08" not in text
    assert "текст" in text and "мусором" in text
