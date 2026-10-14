"""karkas.infra.ingest — извлечение текста и структуры из загруженных файлов.

Поддержка: .docx (python-docx, таблицы и заголовки), .pdf (pypdf), .txt/.md.
Файлы разбираются как данные: макросы не выполняются, .docx читается как XML.
"""

from __future__ import annotations

import io
import re
from pathlib import PurePosixPath
from typing import Any

from karkas.api.errors import UnsupportedFile
from karkas.logging_setup import get_logger
from karkas.settings import settings

logger = get_logger("karkas.ingest")

ALLOWED_MIME = {
    "application/pdf": ".pdf",
    "text/plain": ".txt",
    "text/markdown": ".md",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
}

CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


def detect_mime(filename: str, data: bytes) -> str:
    """Определить MIME по расширению и сигнатуре (проверка подделки типа — Д1 §13)."""
    ext = PurePosixPath(filename).suffix.lower()
    by_ext = {v: k for k, v in ALLOWED_MIME.items()}.get(ext)
    if by_ext is None:
        raise UnsupportedFile(f"Тип {ext or 'без расширения'} не поддерживается (.docx/.pdf/.txt/.md)")

    if ext == ".pdf" and not data.startswith(b"%PDF"):
        raise UnsupportedFile("Файл с расширением .pdf не является PDF")
    if ext == ".docx" and not data.startswith(b"PK"):
        raise UnsupportedFile("Файл с расширением .docx повреждён (ожидался ZIP-контейнер)")
    if ext in (".txt", ".md") and b"\x00" in data[:1024]:
        raise UnsupportedFile("Текстовый файл содержит бинарные данные")
    return by_ext


def verify_size(data: bytes) -> None:
    from karkas.api.errors import FileTooLarge

    if len(data) == 0 or len(data) > settings.file_max_bytes:
        raise FileTooLarge(f"Размер файла должен быть от 1 байта до {settings.file_max_bytes // (1024 * 1024)} МБ")


def _clean(text: str) -> str:
    return CONTROL_CHARS.sub(" ", text).strip()


def extract_docx(data: bytes) -> tuple[str, dict[str, Any]]:
    """Текст, таблицы и заголовки .docx (структура учебных планов сохраняется таблицами)."""
    from docx import Document  # python-docx

    document = Document(io.BytesIO(data))
    parts: list[str] = []
    headings: list[str] = []
    for paragraph in document.paragraphs:
        text = _clean(paragraph.text)
        if not text:
            continue
        style = (paragraph.style.name or "").lower()
        if style.startswith("heading") or style.startswith("заголов"):
            headings.append(text)
            parts.append(f"\n## {text}\n")
        else:
            parts.append(text)
    tables: list[list[list[str]]] = []
    for table in document.tables:
        grid = [[_clean(cell.text) for cell in row.cells] for row in table.rows]
        tables.append(grid)
        parts.append("\n[Таблица]")
        for row in grid:
            parts.append(" | ".join(row))
    meta = {
        "paragraphs": len(document.paragraphs),
        "tables": len(tables),
        "headings": headings,
        "chars": sum(len(p) for p in parts),
    }
    return "\n".join(parts), meta


def extract_pdf(data: bytes) -> tuple[str, dict[str, Any]]:
    """Текст .pdf через pypdf. Скан без текстового слоя → meta.ocr_required."""
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages: list[str] = []
    for page_number, page in enumerate(reader.pages, start=1):
        try:
            pages.append(_clean(page.extract_text() or ""))
        except Exception:
            # Ошибка извлечения страницы не должна ронять загрузку файла,
            # но и молчать нельзя — фиксируем в лог для разбора.
            logger.warning("ingest.pdf_page_failed", page=page_number)
            pages.append("")
    text = "\n\n".join(pages)
    meta = {
        "pages": len(reader.pages),
        "chars": len(text),
        "ocr_required": len(text.strip()) == 0,
    }
    return text, meta


def extract_txt(data: bytes) -> tuple[str, dict[str, Any]]:
    for encoding in ("utf-8", "cp1251"):
        try:
            text = _clean(data.decode(encoding))
            return text, {"chars": len(text), "encoding": encoding, "lines": text.count("\n") + 1}
        except UnicodeDecodeError:
            continue
    raise UnsupportedFile("Не удалось определить кодировку текстового файла")


def extract(filename: str, data: bytes) -> tuple[str, str, dict[str, Any]]:
    """Полный ingest: (mime, извлечённый текст, мета)."""
    verify_size(data)
    mime = detect_mime(filename, data)
    if mime.endswith("wordprocessingml.document"):
        text, meta = extract_docx(data)
    elif mime == "application/pdf":
        text, meta = extract_pdf(data)
    else:
        text, meta = extract_txt(data)
    if len(text) > 2_000_000:
        text = text[:2_000_000]
        meta["truncated"] = True
    return mime, text, meta
