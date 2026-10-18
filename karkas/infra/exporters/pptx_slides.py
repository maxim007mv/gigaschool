"""karkas.infra.exporters.pptx_slides — выгрузка слайдов в .pptx (Д4 §5.12).

Рендер **только** структурированного выхода S13 (GigaChat): заголовки, пункты,
заметки докладчика и выбор макета слайда (`layout`). Код не дописывает
содержательные слайды из компетенций, тестов или планов лекций — это задача ИИ.

Если слайдов нет — минимальный титул с пометкой об отсутствии генерации.
"""

from __future__ import annotations

import io
from typing import Any, Iterable, Sequence

from pptx import Presentation
from pptx.presentation import Presentation as PptxPresentation
from pptx.util import Inches

from karkas.domain.kit import ContentBlock, KitV1, Module

LAYOUT_TITLE = 0
LAYOUT_TITLE_CONTENT = 1
BODY_PLACEHOLDER_IDX = 1
TEXTBOX_POSITION = (Inches(1), Inches(2), Inches(8), Inches(4))
MAX_CONTENT_SLIDES = 12
SLIDES_MISSING_TITLE = "[Слайды не сформированы моделью]"


def _selected_modules(kit: KitV1, module_refs: list[str] | None) -> list[Module]:
    if module_refs is None:
        refs = {block.module_ref for block in kit.content.slides if block.module_ref}
    else:
        refs = set(module_refs)
    return [module for module in kit.framework.modules if module.ref in refs]


def _lines(value: Any) -> list[str]:
    if isinstance(value, list):
        return [str(item) for item in value if str(item).strip()]
    if isinstance(value, str):
        return [line.strip() for line in value.splitlines() if line.strip()]
    return []


def _layout_index(slide: dict[str, Any]) -> int:
    layout = str(slide.get("layout") or slide.get("layout_kind") or "").strip().lower()
    if layout in ("title", "title_only", "титул"):
        return LAYOUT_TITLE
    return LAYOUT_TITLE_CONTENT


def _slide_payloads(blocks: Iterable[ContentBlock]) -> list[tuple[str, list[str], str, int]]:
    """(title, bullets, speaker_notes, layout_index) из body.slides S13."""
    payloads: list[tuple[str, list[str], str, int]] = []
    for block in blocks:
        raw_slides = block.body.get("slides")
        if isinstance(raw_slides, list) and raw_slides:
            for item in raw_slides[:MAX_CONTENT_SLIDES]:
                if isinstance(item, dict):
                    title = str(item.get("title") or "").strip()
                    if not title:
                        continue
                    bullets = (
                        _lines(item.get("bullets"))
                        or _lines(item.get("lines"))
                        or _lines(item.get("text"))
                    )
                    notes = str(item.get("speaker_notes") or item.get("notes") or "").strip()
                    payloads.append((title, bullets, notes, _layout_index(item)))
            continue
        title = str(block.body.get("title") or block.ref).strip()
        bullets = _lines(block.body.get("bullets")) or _lines(block.body)
        notes = str(block.body.get("notes") or "").strip()
        if title:
            payloads.append((title, bullets, notes, LAYOUT_TITLE_CONTENT))
    return payloads[:MAX_CONTENT_SLIDES]


def _add_slide(
    presentation: PptxPresentation,
    layout_idx: int,
    title: str,
    bullets: Sequence[str],
    notes: str | None = None,
) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[layout_idx])
    if slide.shapes.title is not None:
        slide.shapes.title.text = title
    if layout_idx == LAYOUT_TITLE:
        if bullets and len(slide.placeholders) > 1:
            slide.placeholders[1].text = "\n".join(bullets)
    elif layout_idx == LAYOUT_TITLE_CONTENT:
        body = next(
            (shape for shape in slide.placeholders if shape.placeholder_format.idx == BODY_PLACEHOLDER_IDX),
            None,
        )
        if body is not None:
            frame = body.text_frame
            for index, bullet in enumerate(bullets):
                paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                paragraph.text = bullet
        elif bullets:
            textbox = slide.shapes.add_textbox(*TEXTBOX_POSITION)
            frame = textbox.text_frame
            for index, bullet in enumerate(bullets):
                paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                paragraph.text = bullet
    if notes:
        slide.notes_slide.notes_text_frame.text = notes


def _add_missing_slide(presentation: PptxPresentation, kit: KitV1) -> None:
    _add_slide(
        presentation,
        LAYOUT_TITLE,
        kit.program.title,
        [SLIDES_MISSING_TITLE],
        notes="Сгенерируйте слайды навыком S13 (GigaChat) перед выгрузкой.",
    )


def export_pptx_slides(kit: KitV1, module_refs: list[str] | None = None) -> bytes:
    """Собрать .pptx из content.slides (выход S13), без синтетических слайдов."""
    presentation = Presentation()
    modules = _selected_modules(kit, module_refs)
    if not modules:
        _add_missing_slide(presentation, kit)
    for module in modules:
        blocks = [block for block in kit.content.slides if block.module_ref == module.ref]
        payloads = _slide_payloads(blocks)
        if not payloads:
            _add_slide(
                presentation,
                LAYOUT_TITLE,
                f"М{module.number} {module.title}",
                [SLIDES_MISSING_TITLE],
                notes=f"Нет блока SL-* для модуля {module.ref}",
            )
            continue
        for title, bullets, notes, layout_idx in payloads:
            _add_slide(presentation, layout_idx, title, bullets, notes)
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()
