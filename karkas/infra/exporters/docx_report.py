"""karkas.infra.exporters.docx_report — выгрузка отчёта самопроверки в .docx (Д4 §5.8, §5.12).

Структура документа:

- сводка: результат (``passed``), длительность, версия, таблица ``summary``;
- замечания: код правила, серьёзность, объект, сообщение, подсказка;
- допущения: из отчёта и из ``kit.assumptions`` (объединяются без повторов);
- найденные ошибки образца (``sample_findings``);
- разрывы (``gaps``).

Если отчёт не передан и у комплекта нет ``kit.report``, формируется минимальный
корректный документ с пометкой «Отчёт не сформирован» — функция не падает.
"""

from __future__ import annotations

import io
import json
from typing import Any, Sequence

from docx import Document
from docx.document import Document as DocxDocument
from docx.shared import Pt

from karkas.domain.kit import KitV1, Report, ReportIssue

SEVERITY_LABELS: dict[str, str] = {
    "error": "ошибка",
    "warning": "предупреждение",
    "info": "информация",
}
ISSUE_COLUMNS = ("Правило", "Серьёзность", "Объект", "Сообщение", "Подсказка")
MISSING_REPORT_TEXT = "Отчёт не сформирован: запустите самопроверку комплекта (S19)."
PASSED_TEXT = "Проверка пройдена"
FAILED_TEXT = "Проверка не пройдена"
NORMAL_FONT_SIZE_PT = 12


def _prepare_styles(document: DocxDocument) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(NORMAL_FONT_SIZE_PT)


def _para(document: DocxDocument, text: str, *, bold: bool = False) -> None:
    if not text:
        return
    paragraph = document.add_paragraph()
    paragraph.add_run(text).bold = bold


def _set_cell(cell: Any, value: str, *, bold: bool = False) -> None:
    cell.text = str(value)
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = bold


def _summary_value(value: Any) -> str:
    """Значение ``summary``: вложенные структуры — компактным JSON."""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def _add_summary_table(document: DocxDocument, summary: dict[str, Any]) -> None:
    if not summary:
        return
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    _set_cell(table.rows[0].cells[0], "Показатель", bold=True)
    _set_cell(table.rows[0].cells[1], "Значение", bold=True)
    for key, value in summary.items():
        cells = table.add_row().cells
        _set_cell(cells[0], str(key))
        _set_cell(cells[1], _summary_value(value))


def _add_issues_table(document: DocxDocument, issues: Sequence[ReportIssue]) -> None:
    if not issues:
        _para(document, "Замечаний нет.")
        return
    table = document.add_table(rows=1, cols=len(ISSUE_COLUMNS))
    table.style = "Table Grid"
    for index, title in enumerate(ISSUE_COLUMNS):
        _set_cell(table.rows[0].cells[index], title, bold=True)
    for issue in issues:
        cells = table.add_row().cells
        values = [
            issue.rule_code,
            SEVERITY_LABELS.get(issue.severity, issue.severity),
            issue.entity_ref or "",
            issue.message,
            _summary_value(issue.fix_hint) if issue.fix_hint else "",
        ]
        for index, value in enumerate(values):
            _set_cell(cells[index], value)


def _add_bullets(document: DocxDocument, items: Sequence[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def _merged_assumptions(kit: KitV1, report: Report) -> list[str]:
    """Допущения отчёта и комплекта без дублей (порядок сохраняется)."""
    merged: list[str] = []
    for text in [*report.assumptions, *(assumption.text for assumption in kit.assumptions)]:
        if text and text not in merged:
            merged.append(text)
    return merged


def _build_document(kit: KitV1, report: Report) -> bytes:
    document = Document()
    _prepare_styles(document)
    document.add_heading("Отчёт самопроверки комплекта", level=1)
    _para(document, kit.program.title)
    result = PASSED_TEXT if report.passed else FAILED_TEXT
    errors = len(report.errors)
    warnings = len([issue for issue in report.issues if issue.severity == "warning"])
    _para(document, f"Результат: {result}. Ошибок: {errors}, предупреждений: {warnings}.", bold=True)
    version = report.version if report.version is not None else "—"
    _para(document, f"Длительность проверки: {report.duration_ms} мс; версия комплекта: {version}.")

    document.add_heading("Сводные показатели", level=2)
    _add_summary_table(document, report.summary)

    document.add_heading("Замечания", level=2)
    _add_issues_table(document, report.issues)

    document.add_heading("Допущения", level=2)
    assumptions = _merged_assumptions(kit, report)
    if assumptions:
        _add_bullets(document, assumptions)
    else:
        _para(document, "Допущения не зафиксированы.")

    document.add_heading("Найденные ошибки образца", level=2)
    if report.sample_findings:
        _add_bullets(document, report.sample_findings)
    else:
        _para(document, "Ошибки образца не обнаружены.")

    document.add_heading("Разрывы (gaps)", level=2)
    if report.gaps:
        _add_bullets(document, report.gaps)
    else:
        _para(document, "Разрывов между запросом и реалистичным результатом не зафиксировано.")

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _build_empty_document(kit: KitV1) -> bytes:
    document = Document()
    _prepare_styles(document)
    document.add_heading("Отчёт самопроверки комплекта", level=1)
    _para(document, kit.program.title)
    _para(document, MISSING_REPORT_TEXT)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def export_docx_report(kit: KitV1, report: Report | None = None) -> bytes:
    """Экспортировать отчёт самопроверки в .docx.

    Источник отчёта: аргумент ``report``, иначе ``kit.report``. Если отчёта нет —
    возвращается минимальный документ с пометкой «Отчёт не сформирован».
    """
    source = report or kit.report
    if source is None:
        return _build_empty_document(kit)
    return _build_document(kit, source)
