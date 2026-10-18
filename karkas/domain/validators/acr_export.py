"""V-ACR-01: структура выгружаемого .docx совпадает с эталоном АЦР."""

from __future__ import annotations

import io

from docx import Document

from karkas.domain.acr_reference import expected_acr_headings
from karkas.domain.kit import KitV1
from typing import TYPE_CHECKING

from karkas.domain.validators.base import GroupResult, issue, resolve_stage
from karkas.infra.exporters.docx_program import export_docx_program

if TYPE_CHECKING:
    from karkas.domain.validators import ValidationContext


def _docx_plain_text(data: bytes) -> str:
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"не удалось разобрать .docx: {exc}") from exc
    parts: list[str] = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def validate(ctx: ValidationContext) -> GroupResult:
    """V-ACR-01: только на стадии полного комплекта (есть содержание)."""
    if resolve_stage(ctx) != "kit":
        return GroupResult(skipped=["V-ACR-01"], meta={"stage": resolve_stage(ctx)})
    return validate_acr_export(ctx)


def validate_acr_export(ctx: ValidationContext) -> GroupResult:
    kit: KitV1 = ctx.kit
    headings = expected_acr_headings(kit.program.kind, kit.program.form)
    try:
        blob = export_docx_program(kit)
    except Exception as exc:  # noqa: BLE001 — отчёт методисту, не падение API
        return GroupResult(
            issues=[
                issue(
                    "V-ACR-01",
                    "error",
                    f"Не удалось собрать .docx для сверки с АЦР: {exc}",
                    entity_ref="acr_export",
                )
            ],
            meta={"kind": kit.program.kind, "form": kit.program.form},
        )
    text = _docx_plain_text(blob)
    missing = [h for h in headings if h not in text]
    issues = []
    if missing:
        issues.append(
            issue(
                "V-ACR-01",
                "error",
                "В .docx нет разделов эталона АЦР: " + "; ".join(missing[:8])
                + ("…" if len(missing) > 8 else ""),
                entity_ref="acr_export",
                fix_hint={"missing_headings": missing},
            )
        )
    ref = "data/templates/"
    return GroupResult(
        issues=issues,
        meta={
            "expected_sections": len(headings),
            "missing_count": len(missing),
            "reference": ref + ("acr_dop_intensive_ai_24h.txt" if kit.program.kind == "dop" else "acr_dpp_pk_staj_bsa_248h.txt"),
        },
    )
