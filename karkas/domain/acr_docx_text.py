"""Извлечение текста из .docx для проверки структуры АЦР (без сборки документа)."""

from __future__ import annotations

import io

from docx import Document

from karkas.domain.acr_reference import expected_acr_headings


def docx_plain_text(data: bytes) -> str:
    """Текст параграфов и ячеек таблиц."""
    document = Document(io.BytesIO(data))
    parts: list[str] = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def missing_acr_headings(text: str, kind: str, form: str | None) -> list[str]:
    headings = expected_acr_headings(kind, form)
    return [h for h in headings if h not in text]
